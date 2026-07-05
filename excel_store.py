import io
import os
import re
import shutil
import tempfile
from datetime import date, datetime

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import currency
import settings
from config import CATEGORIES, CURRENCIES, CURRENCY_LABELS, DATA_DIR, EXCEL_PATH

# Имена листов одинаковы во всех файлах: каждый тур — отдельный .xlsx, поэтому
# уникальность имён листов между турами больше не нужна.
DATA_SHEET = "Расходы"
SUMMARY_SHEET = "Итоги"
WALLET_SHEET = "Кошелёк"

_HEADERS = ["Дата", "Описание", "Сумма", "Валюта", "Категория", "Сумма USD", "ID"]
_DATA_WIDTHS = [12, 34, 12, 10, 14, 12, 8]
_ID_COL = 7  # технический столбец стабильного идентификатора (скрыт)
_WALLET_HEADERS = ["Дата", "Тип", "Валюта", "Сумма", "Примечание"]
_WALLET_WIDTHS = [12, 16, 10, 14, 30]
_WALLET_ID_COL = 6  # скрытый столбец стабильного id движения кошелька
_START_KIND = "Старт"  # вид движения «начальный баланс» кошелька
_START_NOTE = "Начальный баланс"
_ADJUST_KIND = "Коррекция"  # ручная правка остатка владельцем (set_balance)
_HEADER_FILL = PatternFill("solid", fgColor="4472C4")
_HEADER_FONT = Font(bold=True, color="FFFFFF")
_TITLE_FONT = Font(bold=True, size=12)
# Символы, недопустимые в имени файла (Windows/Unix) + управляющие.
_INVALID_FILE_CHARS = re.compile(r'[\\/:*?"<>|\x00-\x1f]')


# --- Пути к файлам туров ----------------------------------------------------


def _active_path() -> str:
    return os.path.join(DATA_DIR, settings.get_active_file())


def active_path() -> str:
    """Публичный путь к файлу активного тура (для выгрузки /excel)."""
    return _active_path()


def _path_for(name: str | None = None) -> str | None:
    """Путь к файлу тура по имени (или активного, если name=None)."""
    file = settings.get_active_file() if name is None else settings.get_profile_file(name)
    return os.path.join(DATA_DIR, file) if file else None


def _safe_filename(name: str, existing: set[str] | None = None) -> str:
    """Имя .xlsx-файла из названия тура: без запрещённых символов, уникальное в DATA_DIR."""
    cleaned = _INVALID_FILE_CHARS.sub(" ", name or "")
    cleaned = re.sub(r"\s+", " ", cleaned).strip().rstrip(". ")
    cleaned = (cleaned[:80].strip() or "Тур")
    if existing is None:
        existing = (
            {f.lower() for f in os.listdir(DATA_DIR)} if os.path.isdir(DATA_DIR) else set()
        )
    candidate = f"{cleaned}.xlsx"
    n = 2
    while candidate.lower() in existing:
        candidate = f"{cleaned} {n}.xlsx"
        n += 1
    return candidate


def _atomic_save(wb: Workbook, path: str) -> None:
    """Сохраняет книгу через временный файл + os.replace, чтобы аварийное
    завершение (сон Мака, перезапуск) не оставило битый .xlsx."""
    folder = os.path.dirname(os.path.abspath(path)) or "."
    os.makedirs(folder, exist_ok=True)
    fd, tmp = tempfile.mkstemp(dir=folder, prefix=".tour_", suffix=".xlsx")
    os.close(fd)
    try:
        wb.save(tmp)
        os.replace(tmp, path)
    except Exception:
        if os.path.exists(tmp):
            os.remove(tmp)
        raise


# --- Построение листов ------------------------------------------------------


def _init_wallet_sheet(ws) -> None:
    for col, header in enumerate(_WALLET_HEADERS, start=1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="center")
    for i, w in enumerate(_WALLET_WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    idc = ws.cell(row=1, column=_WALLET_ID_COL, value="ID")
    idc.fill = _HEADER_FILL
    idc.font = _HEADER_FONT
    idc.alignment = Alignment(horizontal="center")
    ws.column_dimensions[get_column_letter(_WALLET_ID_COL)].hidden = True
    ws.freeze_panes = "A2"


def _init_data_sheet(ws) -> None:
    for col, header in enumerate(_HEADERS, start=1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="center")
    for i, w in enumerate(_DATA_WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.column_dimensions[get_column_letter(_ID_COL)].hidden = True
    ws.freeze_panes = "A2"


def _new_workbook() -> Workbook:
    """Свежий файл тура: листы Расходы / Итоги / Кошелёк."""
    wb = Workbook()
    ws = wb.active
    ws.title = DATA_SHEET
    _init_data_sheet(ws)
    wb.create_sheet(SUMMARY_SHEET)
    _init_wallet_sheet(wb.create_sheet(WALLET_SHEET))
    return wb


def _load(path: str) -> Workbook:
    """Открывает файл тура, добавляя недостающие стандартные листы. Если файла нет —
    создаёт новую книгу (на диск пишет только вызывающий через _atomic_save)."""
    if os.path.exists(path):
        wb = load_workbook(path)
        if DATA_SHEET not in wb.sheetnames:
            _init_data_sheet(wb.create_sheet(DATA_SHEET))
        if SUMMARY_SHEET not in wb.sheetnames:
            wb.create_sheet(SUMMARY_SHEET)
        if WALLET_SHEET not in wb.sheetnames:
            _init_wallet_sheet(wb.create_sheet(WALLET_SHEET))
    else:
        wb = _new_workbook()
    return wb


def _next_wallet_id(ws) -> int:
    """Следующий свободный id движения кошелька (max существующих + 1)."""
    mx = 0
    for row in ws.iter_rows(min_row=2, values_only=True):
        v = row[_WALLET_ID_COL - 1] if len(row) >= _WALLET_ID_COL else None
        if isinstance(v, int) and v > mx:
            mx = v
    return mx + 1


def _find_wallet_row_by_id(ws, movement_id: int) -> int | None:
    for row in range(2, ws.max_row + 1):
        if ws.cell(row=row, column=_WALLET_ID_COL).value == movement_id:
            return row
    return None


def _ensure_wallet_ids(ws) -> bool:
    """Миграция листа «Кошелёк»: гарантирует скрытый заголовок ID и проставляет
    стабильные id движениям без него."""
    changed = False
    if ws.cell(row=1, column=_WALLET_ID_COL).value != "ID":
        cell = ws.cell(row=1, column=_WALLET_ID_COL, value="ID")
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="center")
        ws.column_dimensions[get_column_letter(_WALLET_ID_COL)].hidden = True
        changed = True
    nxt = _next_wallet_id(ws)
    for row in range(2, ws.max_row + 1):
        has_data = ws.cell(row=row, column=2).value is not None
        cur_id = ws.cell(row=row, column=_WALLET_ID_COL).value
        if has_data and not isinstance(cur_id, int):
            ws.cell(row=row, column=_WALLET_ID_COL, value=nxt)
            nxt += 1
            changed = True
    return changed


def _wallet_row_dict(ws, row: int) -> dict | None:
    if row < 2 or row > ws.max_row:
        return None
    kind = ws.cell(row=row, column=2).value
    if kind is None:
        return None
    d = ws.cell(row=row, column=1).value
    return {
        "id": ws.cell(row=row, column=_WALLET_ID_COL).value,
        "date": d.date().isoformat() if isinstance(d, datetime) else (str(d) if d else ""),
        "kind": kind,
        "currency": ws.cell(row=row, column=3).value,
        "amount": ws.cell(row=row, column=4).value,
        "note": ws.cell(row=row, column=5).value,
    }


def _append_movement(
    ws, kind: str, currency: str, amount: float, note: str = "", on_date: date | None = None
) -> None:
    _ensure_wallet_ids(ws)  # гарантируем скрытый столбец id перед записью
    next_row = ws.max_row + 1 if ws.max_row >= 1 else 2
    if next_row < 2:
        next_row = 2
    d = ws.cell(row=next_row, column=1, value=on_date or date.today())
    d.number_format = "yyyy-mm-dd"
    ws.cell(row=next_row, column=2, value=kind)
    ws.cell(row=next_row, column=3, value=currency)
    c = ws.cell(row=next_row, column=4, value=round(float(amount), 2))
    c.number_format = "#,##0.00"
    ws.cell(row=next_row, column=5, value=note)
    ws.cell(row=next_row, column=_WALLET_ID_COL, value=_next_wallet_id(ws))


def _wallet_net_from_ws(ws) -> dict:
    net = {cur: 0.0 for cur in CURRENCIES}
    for row in ws.iter_rows(min_row=2, values_only=True):
        cur, amt = row[2], row[3]
        if cur in net and isinstance(amt, (int, float)):
            net[cur] += amt
    return {cur: round(v, 2) for cur, v in net.items()}


def _next_id(ws) -> int:
    """Следующий свободный идентификатор траты на листе (max существующих + 1)."""
    mx = 0
    for row in ws.iter_rows(min_row=2, values_only=True):
        v = row[_ID_COL - 1] if len(row) >= _ID_COL else None
        if isinstance(v, int) and v > mx:
            mx = v
    return mx + 1


def _find_row_by_id(ws, expense_id: int) -> int | None:
    """Физический номер строки по стабильному id, либо None."""
    for row in range(2, ws.max_row + 1):
        if ws.cell(row=row, column=_ID_COL).value == expense_id:
            return row
    return None


def _ensure_ids(ws) -> bool:
    """Миграция: гарантирует заголовок ID и присваивает id строкам без него."""
    changed = False
    if ws.cell(row=1, column=_ID_COL).value != "ID":
        cell = ws.cell(row=1, column=_ID_COL, value="ID")
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="center")
        ws.column_dimensions[get_column_letter(_ID_COL)].hidden = True
        changed = True
    nxt = _next_id(ws)
    for row in range(2, ws.max_row + 1):
        has_data = ws.cell(row=row, column=3).value is not None
        cur_id = ws.cell(row=row, column=_ID_COL).value
        if has_data and not isinstance(cur_id, int):
            ws.cell(row=row, column=_ID_COL, value=nxt)
            nxt += 1
            changed = True
    return changed


def _parse_date(value: str):
    try:
        return datetime.strptime(value, "%Y-%m-%d").date()
    except (ValueError, TypeError):
        return date.today()


def _coerce_date(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        try:
            return datetime.strptime(value, "%Y-%m-%d").date()
        except ValueError:
            return None
    return None


def _aggregate(ws, since: date | None = None, up_to_id: int | None = None) -> dict:
    """Суммирует строки листа расходов в готовые числа (без формул Excel).
    since — нижняя граница даты включительно (для периодных сводок).
    up_to_id — учитывать только траты с id не больше указанного: так остаток
    в ответе на старую трату не включает записанные позже."""
    per_currency = {cur: 0.0 for cur in CURRENCIES}
    per_category = {cat: 0.0 for cat in CATEGORIES}
    total_usd = 0.0
    count = 0
    for row in ws.iter_rows(min_row=2, values_only=True):
        amount, cur, cat, usd = row[2], row[3], row[4], row[5]
        if amount is None and usd is None:
            continue
        if since is not None:
            d = _coerce_date(row[0])
            if d is None or d < since:
                continue
        if up_to_id is not None:
            rid = row[_ID_COL - 1] if len(row) >= _ID_COL else None
            if isinstance(rid, int) and rid > up_to_id:
                continue
        count += 1
        if cur in per_currency and isinstance(amount, (int, float)):
            per_currency[cur] += amount
        if cat in per_category and isinstance(usd, (int, float)):
            per_category[cat] += usd
        if isinstance(usd, (int, float)):
            total_usd += usd
    return {
        "per_currency": per_currency,
        "per_category": per_category,
        "total_usd": round(total_usd, 2),
        "count": count,
    }


def _rebuild_summary(
    wb: Workbook, data_sheet: str, summary_sheet: str, wallet_sheet: str | None = None
) -> None:
    ws = wb[summary_sheet]
    ws.delete_rows(1, ws.max_row or 1)
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 16

    totals = _aggregate(wb[data_sheet])
    net = None
    if wallet_sheet and wallet_sheet in wb.sheetnames:
        net = _wallet_net_from_ws(wb[wallet_sheet])
    row = 1

    if net is not None:
        ws.cell(row=row, column=1, value="Остаток в кошельке").font = _TITLE_FONT
        row += 1
        for cur in CURRENCIES:
            label = f"{CURRENCY_LABELS.get(cur, cur)} ({cur})"
            remaining = round(net.get(cur, 0.0) - totals["per_currency"][cur], 2)
            ws.cell(row=row, column=1, value=label)
            c = ws.cell(row=row, column=2, value=remaining)
            c.number_format = "#,##0.00"
            row += 1
        row += 1

    ws.cell(row=row, column=1, value="Итого по валютам").font = _TITLE_FONT
    row += 1
    for cur in CURRENCIES:
        label = f"{CURRENCY_LABELS.get(cur, cur)} ({cur})"
        ws.cell(row=row, column=1, value=label)
        c = ws.cell(row=row, column=2, value=round(totals["per_currency"][cur], 2))
        c.number_format = "#,##0.00"
        row += 1

    row += 1
    ws.cell(row=row, column=1, value="Всего в USD").font = _TITLE_FONT
    c = ws.cell(row=row, column=2, value=totals["total_usd"])
    c.number_format = "#,##0.00"
    c.font = _TITLE_FONT
    row += 2

    ws.cell(row=row, column=1, value="Итого по категориям (USD)").font = _TITLE_FONT
    row += 1
    for cat in CATEGORIES:
        ws.cell(row=row, column=1, value=cat)
        c = ws.cell(row=row, column=2, value=round(totals["per_category"][cat], 2))
        c.number_format = "#,##0.00"
        row += 1


def compute_totals(since: date | None = None, up_to_id: int | None = None) -> dict:
    """Считает итоги по активному туру: суммы по валютам, общую в USD и суммы по
    категориям (в USD). since — нижняя граница даты включительно.
    up_to_id — только траты с id не больше указанного (остаток «на момент траты»)."""
    empty = {
        "per_currency": {cur: 0.0 for cur in CURRENCIES},
        "per_category": {cat: 0.0 for cat in CATEGORIES},
        "total_usd": 0.0,
        "count": 0,
    }
    path = _active_path()
    if not os.path.exists(path):
        return empty
    wb = load_workbook(path, data_only=True)
    if DATA_SHEET not in wb.sheetnames:
        return empty
    return _aggregate(wb[DATA_SHEET], since, up_to_id)


def add_expense(entry: dict) -> int:
    """entry: {date, description, amount, currency, category, amount_usd}.
    Пишет в файл активного тура, возвращает стабильный id траты."""
    path = _active_path()
    wb = _load(path)
    ws = wb[DATA_SHEET]
    _ensure_ids(ws)
    next_row = ws.max_row + 1 if ws.max_row >= 1 else 2
    if next_row < 2:
        next_row = 2
    expense_id = _next_id(ws)

    date_cell = ws.cell(row=next_row, column=1, value=_parse_date(entry["date"]))
    date_cell.number_format = "yyyy-mm-dd"
    ws.cell(row=next_row, column=2, value=entry["description"])
    amt = ws.cell(row=next_row, column=3, value=entry["amount"])
    amt.number_format = "#,##0.00"
    ws.cell(row=next_row, column=4, value=entry["currency"])
    ws.cell(row=next_row, column=5, value=entry["category"])
    usd = ws.cell(row=next_row, column=6, value=entry["amount_usd"])
    usd.number_format = "#,##0.00"
    ws.cell(row=next_row, column=_ID_COL, value=expense_id)

    _rebuild_summary(wb, DATA_SHEET, SUMMARY_SHEET, WALLET_SHEET)
    _atomic_save(wb, path)
    return expense_id


def _row_to_entry(ws, row: int) -> dict | None:
    if row < 2 or row > ws.max_row:
        return None
    d, desc, amt, cur, cat, usd = (ws.cell(row=row, column=c).value for c in range(1, 7))
    if amt is None and usd is None:
        return None
    return {
        "date": d.date().isoformat() if isinstance(d, datetime) else str(d),
        "description": desc,
        "amount": amt,
        "currency": cur,
        "category": cat,
        "amount_usd": usd,
        "id": ws.cell(row=row, column=_ID_COL).value,
    }


def get_expense(expense_id: int) -> dict | None:
    path = _active_path()
    if not os.path.exists(path):
        return None
    wb = load_workbook(path)
    if DATA_SHEET not in wb.sheetnames:
        return None
    ws = wb[DATA_SHEET]
    row = _find_row_by_id(ws, expense_id)
    return _row_to_entry(ws, row) if row else None


def last_entry() -> dict | None:
    """Последняя добавленная трата активного тура (для /undo), либо None."""
    path = _active_path()
    if not os.path.exists(path):
        return None
    wb = load_workbook(path)
    if DATA_SHEET not in wb.sheetnames:
        return None
    ws = wb[DATA_SHEET]
    for row in range(ws.max_row, 1, -1):
        if ws.cell(row=row, column=3).value is not None:
            return _row_to_entry(ws, row)
    return None


def update_category(expense_id: int, category: str) -> dict | None:
    """Меняет категорию траты по стабильному id."""
    path = _active_path()
    wb = _load(path)
    ws = wb[DATA_SHEET]
    row = _find_row_by_id(ws, expense_id)
    entry = _row_to_entry(ws, row) if row else None
    if entry is None:
        return None
    ws.cell(row=row, column=5, value=category)
    entry["category"] = category
    _rebuild_summary(wb, DATA_SHEET, SUMMARY_SHEET, WALLET_SHEET)
    _atomic_save(wb, path)
    return entry


def update_amount(
    expense_id: int, amount: float, cur: str | None = None
) -> dict | None:
    """Меняет сумму (и при желании валюту) траты по id, пересчитывает USD."""
    path = _active_path()
    wb = _load(path)
    ws = wb[DATA_SHEET]
    row = _find_row_by_id(ws, expense_id)
    entry = _row_to_entry(ws, row) if row else None
    if entry is None:
        return None
    amount = round(float(amount), 2)
    new_cur = (cur or entry["currency"] or "USD").upper()
    if new_cur not in CURRENCIES:
        new_cur = entry["currency"]
    usd = currency.to_usd(amount, new_cur)
    ws.cell(row=row, column=3, value=amount).number_format = "#,##0.00"
    ws.cell(row=row, column=4, value=new_cur)
    ws.cell(row=row, column=6, value=usd).number_format = "#,##0.00"
    entry.update(amount=amount, currency=new_cur, amount_usd=usd)
    _rebuild_summary(wb, DATA_SHEET, SUMMARY_SHEET, WALLET_SHEET)
    _atomic_save(wb, path)
    return entry


def update_description(expense_id: int, description: str) -> dict | None:
    """Меняет описание траты по id."""
    path = _active_path()
    wb = _load(path)
    ws = wb[DATA_SHEET]
    row = _find_row_by_id(ws, expense_id)
    entry = _row_to_entry(ws, row) if row else None
    if entry is None:
        return None
    description = (description or "").strip() or entry["description"]
    ws.cell(row=row, column=2, value=description)
    entry["description"] = description
    _rebuild_summary(wb, DATA_SHEET, SUMMARY_SHEET, WALLET_SHEET)
    _atomic_save(wb, path)
    return entry


def delete_expense(expense_id: int) -> dict | None:
    """Удаляет трату по стабильному id (id остальных строк не меняются)."""
    path = _active_path()
    wb = _load(path)
    ws = wb[DATA_SHEET]
    row = _find_row_by_id(ws, expense_id)
    entry = _row_to_entry(ws, row) if row else None
    if entry is None:
        return None
    ws.delete_rows(row, 1)
    _rebuild_summary(wb, DATA_SHEET, SUMMARY_SHEET, WALLET_SHEET)
    _atomic_save(wb, path)
    return entry


# --- Туры (профили) ---------------------------------------------------------


def create_profile(name: str) -> dict:
    """Создаёт новый тур: отдельный .xlsx-файл с листами Расходы/Итоги/Кошелёк
    и запись в настройках."""
    name = name.strip()
    filename = _safe_filename(name)
    path = os.path.join(DATA_DIR, filename)
    wb = _new_workbook()
    _rebuild_summary(wb, DATA_SHEET, SUMMARY_SHEET, WALLET_SHEET)
    _atomic_save(wb, path)
    settings.register_profile(name, filename)
    return {"name": name, "file": filename}


def delete_profile(name: str) -> dict | None:
    """Удаляет тур: убирает запись из настроек и переносит его .xlsx в бэкап
    (файл не стираем — на случай ошибки его можно вернуть). Последний тур удалить
    нельзя. Возвращает {name, file, backup} или None, если удалить не получилось."""
    file = settings.get_profile_file(name)
    if not settings.delete_profile(name):
        return None
    backup = None
    if file:
        path = os.path.join(DATA_DIR, file)
        if os.path.exists(path):
            stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            base = os.path.splitext(file)[0]
            backup = f"{base}_deleted_{stamp}.xlsx"
            dst = os.path.join(DATA_DIR, backup)
            try:
                os.replace(path, dst)
            except OSError:
                shutil.copy2(path, dst)
                os.remove(path)
    return {"name": name, "file": file, "backup": backup}


def _normalize_imported(wb: Workbook) -> None:
    """Приводит присланную книгу к структуре файла тура: первый лист считаем
    расходами, гарантируем Итоги/Кошелёк и стабильные id."""
    if DATA_SHEET not in wb.sheetnames:
        wb[wb.sheetnames[0]].title = DATA_SHEET
    _ensure_ids(wb[DATA_SHEET])
    if SUMMARY_SHEET not in wb.sheetnames:
        wb.create_sheet(SUMMARY_SHEET)
    if WALLET_SHEET not in wb.sheetnames:
        _init_wallet_sheet(wb.create_sheet(WALLET_SHEET))


def import_profile_from_upload(filename: str, data: bytes) -> tuple[str, bool]:
    """Импорт тура из присланного .xlsx. Имя файла → название тура: создаёт новый
    тур или заменяет файл существующего (с бэкапом), делает его активным.
    Возвращает (имя тура, replaced). Бросает исключение, если это не .xlsx."""
    wb = load_workbook(io.BytesIO(data))  # битый/не-xlsx → исключение
    stem = os.path.splitext(os.path.basename(filename or ""))[0]
    stem = re.sub(r"\s+", " ", _INVALID_FILE_CHARS.sub(" ", stem)).strip()
    name = stem or "Импортированный тур"

    _normalize_imported(wb)

    replaced = settings.profile_exists(name)
    if replaced:
        target = settings.get_profile_file(name) or _safe_filename(name)
        path = os.path.join(DATA_DIR, target)
        if os.path.exists(path):
            stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            base = os.path.splitext(target)[0]
            shutil.copy2(path, os.path.join(DATA_DIR, f"{base}_backup_{stamp}.xlsx"))
    else:
        target = _safe_filename(name)
        path = os.path.join(DATA_DIR, target)

    _rebuild_summary(wb, DATA_SHEET, SUMMARY_SHEET, WALLET_SHEET)
    _atomic_save(wb, path)
    settings.register_profile(name, target)
    settings.set_active_profile(name)
    return name, replaced


# --- Кошелёк ----------------------------------------------------------------
# Лист «Кошелёк» — единственный источник истины по остатку. Каждая строка —
# движение (Старт/Пополнение/Обмен) со знаком: поступления «+», выдачи «−».
# Чистый остаток = сумма движений; доступно = остаток − траты по валюте.


def wallet_net(name: str | None = None) -> dict:
    """Чистый остаток кошелька тура по валютам (поступления + обмены − выдачи)."""
    empty = {cur: 0.0 for cur in CURRENCIES}
    path = _path_for(name)
    if not path or not os.path.exists(path):
        return empty
    wb = load_workbook(path, data_only=True)
    if WALLET_SHEET not in wb.sheetnames:
        return empty
    return _wallet_net_from_ws(wb[WALLET_SHEET])


def add_movement(kind: str, currency: str, amount: float, note: str = "") -> dict:
    """Добавляет одно движение в кошелёк активного тура, возвращает остаток."""
    path = _active_path()
    wb = _load(path)
    ws = wb[WALLET_SHEET]
    _append_movement(ws, kind, currency, amount, note)
    _rebuild_summary(wb, DATA_SHEET, SUMMARY_SHEET, WALLET_SHEET)
    _atomic_save(wb, path)
    return _wallet_net_from_ws(ws)


def add_exchange(
    out_cur: str,
    out_amt: float,
    in_cur: str,
    in_amt: float,
    kind: str = "Обмен",
) -> dict:
    """Записывает обмен: списывает out_cur и зачисляет in_cur. Возвращает остаток."""
    path = _active_path()
    wb = _load(path)
    ws = wb[WALLET_SHEET]
    _append_movement(ws, kind, out_cur, -abs(float(out_amt)), f"→ {in_amt:,.2f} {in_cur}")
    _append_movement(ws, kind, in_cur, abs(float(in_amt)), f"← {out_amt:,.2f} {out_cur}")
    _rebuild_summary(wb, DATA_SHEET, SUMMARY_SHEET, WALLET_SHEET)
    _atomic_save(wb, path)
    return _wallet_net_from_ws(ws)


def list_movements(name: str | None = None) -> list[dict]:
    """Движения кошелька тура в порядке записи: id/дата/тип/валюта/сумма/примечание."""
    path = _path_for(name)
    if not path or not os.path.exists(path):
        return []
    wb = load_workbook(path, data_only=True)
    if WALLET_SHEET not in wb.sheetnames:
        return []
    ws = wb[WALLET_SHEET]
    out = []
    for row in range(2, ws.max_row + 1):
        m = _wallet_row_dict(ws, row)
        if m is not None:
            out.append(m)
    return out


def delete_movement(movement_id: int) -> list[dict] | None:
    """Удаляет движение кошелька активного тура по стабильному id. Обмен/Возврат
    удаляется обеими ногами (списание + зачисление), чтобы остаток не разъехался.
    Возвращает список удалённых движений (1 или 2) либо None, если id не найден."""
    path = _active_path()
    wb = _load(path)
    ws = wb[WALLET_SHEET]
    _ensure_wallet_ids(ws)
    row = _find_wallet_row_by_id(ws, movement_id)
    if row is None:
        return None
    target = _wallet_row_dict(ws, row)
    rows = {row}
    deleted = [target]
    note = target.get("note") or ""
    # Обмен/Возврат пишутся двумя соседними строками с зеркальными стрелками (→/←):
    # удаляем парную ногу той же операции, иначе кошелёк станет несбалансированным.
    if target["kind"] in ("Обмен", "Возврат") and (note.startswith("→") or note.startswith("←")):
        mirror = "←" if note.startswith("→") else "→"
        for sib in (row + 1, row - 1):
            if 2 <= sib <= ws.max_row and sib not in rows:
                s = _wallet_row_dict(ws, sib)
                if s and s["kind"] == target["kind"] and (s.get("note") or "").startswith(mirror):
                    rows.add(sib)
                    deleted.append(s)
                    break
    for r in sorted(rows, reverse=True):
        ws.delete_rows(r, 1)
    _rebuild_summary(wb, DATA_SHEET, SUMMARY_SHEET, WALLET_SHEET)
    _atomic_save(wb, path)
    return deleted


def get_start_balance(name: str | None = None) -> dict:
    """Начальный баланс тура по валютам — сумма строк «Старт» в кошельке."""
    result = {cur: 0.0 for cur in CURRENCIES}
    path = _path_for(name)
    if not path or not os.path.exists(path):
        return result
    wb = load_workbook(path, data_only=True)
    if WALLET_SHEET not in wb.sheetnames:
        return result
    for row in wb[WALLET_SHEET].iter_rows(min_row=2, values_only=True):
        kind, cur, amt = row[1], row[2], row[3]
        if kind == _START_KIND and cur in result and isinstance(amt, (int, float)):
            result[cur] += amt
    return {cur: round(v, 2) for cur, v in result.items()}


def set_start_balance(pairs: list[tuple[float, str]]) -> dict:
    """Переписывает начальный баланс активного тура: удаляет прежние строки «Старт»
    и записывает новые. Пополнения и обмены не трогаются. Пустой список обнуляет
    начальный баланс. Дата прежнего «Старта» сохраняется, чтобы правка не «сдвигала»
    начальный баланс на сегодня. Возвращает остаток кошелька."""
    path = _active_path()
    wb = _load(path)
    ws = wb[WALLET_SHEET]

    start_date = None
    for row in range(ws.max_row, 1, -1):
        if ws.cell(row=row, column=2).value == _START_KIND:
            d = _coerce_date(ws.cell(row=row, column=1).value)
            if d is not None and (start_date is None or d < start_date):
                start_date = d
            ws.delete_rows(row, 1)

    # Складываем одинаковые валюты, чтобы повторные правки не плодили строки.
    merged: dict[str, float] = {}
    for amount, cur in pairs:
        if cur in CURRENCIES:
            merged[cur] = merged.get(cur, 0.0) + float(amount)
    for cur in CURRENCIES:
        if merged.get(cur):
            _append_movement(ws, _START_KIND, cur, merged[cur], _START_NOTE, start_date)

    _rebuild_summary(wb, DATA_SHEET, SUMMARY_SHEET, WALLET_SHEET)
    _atomic_save(wb, path)
    return _wallet_net_from_ws(ws)


def set_balance(pairs: list[tuple[float, str]]) -> dict:
    """Делает текущий остаток (кошелёк − траты) по каждой валюте ровно равным
    заданному, дописывая одно движение «Коррекция» на разницу. Прежние движения и
    траты не трогаются — правка прозрачна и её видно/можно удалить в истории.
    Возвращает чистый остаток кошелька (net) после правки."""
    path = _active_path()
    wb = _load(path)
    ws = wb[WALLET_SHEET]
    net = _wallet_net_from_ws(ws)
    totals = _aggregate(wb[DATA_SHEET])

    merged: dict[str, float] = {}
    for amount, cur in pairs:
        if cur in CURRENCIES:
            merged[cur] = float(amount)  # последнее значение по валюте побеждает
    for cur, target in merged.items():
        remaining = round(net.get(cur, 0.0) - totals["per_currency"].get(cur, 0.0), 2)
        delta = round(target - remaining, 2)
        if delta:
            _append_movement(ws, _ADJUST_KIND, cur, delta, f"Остаток → {target:,.2f} {cur}")

    _rebuild_summary(wb, DATA_SHEET, SUMMARY_SHEET, WALLET_SHEET)
    _atomic_save(wb, path)
    return _wallet_net_from_ws(ws)


# --- Миграции ---------------------------------------------------------------


def _copy_sheet_rows(src_wb: Workbook, sheet_name: str | None, dst_ws) -> None:
    """Копирует строки данных (со 2-й) из листа исходной книги в целевой лист,
    сохраняя значения и числовые форматы."""
    if not sheet_name or sheet_name not in src_wb.sheetnames:
        return
    src = src_wb[sheet_name]
    r = dst_ws.max_row
    for row in src.iter_rows(min_row=2):
        if all(c.value is None for c in row):
            continue
        r += 1
        for c in row:
            nc = dst_ws.cell(row=r, column=c.column, value=c.value)
            if c.has_style:
                nc.number_format = c.number_format


def migrate_to_files() -> bool:
    """Разовая миграция со старой схемы (все туры в одном файле, листы на тур) на
    схему «файл на тур». Идемпотентна: ничего не делает, если у всех туров уже есть
    `file`. Старый общий файл бэкапится и выводится из обращения."""
    profiles = settings.all_profiles()
    legacy = {
        n: s for n, s in profiles.items()
        if isinstance(s, dict) and "file" not in s
    }
    if not legacy:
        return False

    legacy_wb = load_workbook(EXCEL_PATH) if os.path.exists(EXCEL_PATH) else None
    existing = {f.lower() for f in os.listdir(DATA_DIR)} if os.path.isdir(DATA_DIR) else set()

    for name, sheets in legacy.items():
        filename = _safe_filename(name, existing)
        existing.add(filename.lower())
        path = os.path.join(DATA_DIR, filename)
        wb = _new_workbook()
        if legacy_wb is not None:
            _copy_sheet_rows(legacy_wb, sheets.get("data"), wb[DATA_SHEET])
            _copy_sheet_rows(legacy_wb, sheets.get("wallet_sheet"), wb[WALLET_SHEET])
        _ensure_ids(wb[DATA_SHEET])
        _rebuild_summary(wb, DATA_SHEET, SUMMARY_SHEET, WALLET_SHEET)
        _atomic_save(wb, path)
        settings.register_profile(name, filename)  # перетирает старую запись на {file}

    # Старый общий файл сохраняем как бэкап и больше не используем.
    if legacy_wb is not None and os.path.exists(EXCEL_PATH):
        folder = os.path.dirname(os.path.abspath(EXCEL_PATH)) or "."
        stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        backup = os.path.join(folder, f"expenses_legacy_backup_{stamp}.xlsx")
        try:
            os.replace(EXCEL_PATH, backup)
        except OSError:
            shutil.copy2(EXCEL_PATH, backup)
    return True


def ensure_ids_setup() -> None:
    """Миграция: проставляет стабильные ID тратам во всех файлах туров."""
    for name in settings.list_profiles():
        path = _path_for(name)
        if not path or not os.path.exists(path):
            continue
        wb = load_workbook(path)
        if DATA_SHEET in wb.sheetnames and _ensure_ids(wb[DATA_SHEET]):
            _atomic_save(wb, path)


def ensure_wallet_ids_setup() -> None:
    """Миграция: проставляет стабильные id движениям кошелька во всех файлах туров,
    чтобы их можно было удалять по id (см. delete_movement)."""
    for name in settings.list_profiles():
        path = _path_for(name)
        if not path or not os.path.exists(path):
            continue
        wb = load_workbook(path)
        if WALLET_SHEET in wb.sheetnames and _ensure_wallet_ids(wb[WALLET_SHEET]):
            _atomic_save(wb, path)


def ensure_wallet_setup() -> None:
    """Переносит совсем старый остаток кошелька из settings.json (если ещё остался)
    в строки «Старт» соответствующего файла тура."""
    for name in settings.list_profiles():
        legacy = settings.get_legacy_wallet(name)
        if not legacy:
            continue
        path = _path_for(name)
        if not path:
            continue
        wb = _load(path)
        ws = wb[WALLET_SHEET]
        if ws.max_row < 2:
            for cur in CURRENCIES:
                amt = legacy.get(cur)
                if isinstance(amt, (int, float)) and amt:
                    _append_movement(ws, _START_KIND, cur, amt, _START_NOTE)
            _rebuild_summary(wb, DATA_SHEET, SUMMARY_SHEET, WALLET_SHEET)
            _atomic_save(wb, path)
        settings.clear_legacy_wallet(name)
