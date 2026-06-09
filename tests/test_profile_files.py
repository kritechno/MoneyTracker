import glob
import importlib
import io
import json
import os
import shutil
import tempfile
import unittest

from openpyxl import Workbook


def _reload_stack():
    """Перезагружает config → settings → excel_store, чтобы подхватить пути из env."""
    import config
    import settings
    import excel_store

    importlib.reload(config)
    importlib.reload(settings)
    importlib.reload(excel_store)
    return settings, excel_store


def _upload_bytes(rows):
    """Готовит байты .xlsx с листом «Расходы» и строками трат (для импорта тура).
    rows: список (описание, сумма, валюта, категория, usd, id)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Расходы"
    ws.append(["Дата", "Описание", "Сумма", "Валюта", "Категория", "Сумма USD", "ID"])
    for desc, amt, cur, cat, usd, rid in rows:
        ws.append(["2026-06-08", desc, amt, cur, cat, usd, rid])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


class ProfileFilesTest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.mkdtemp()
        os.environ["EXCEL_PATH"] = os.path.join(self.tmp, "expenses.xlsx")
        os.environ["SETTINGS_PATH"] = os.path.join(self.tmp, "settings.json")
        self.settings, self.store = _reload_stack()
        self.assertTrue(self.store.DATA_DIR.startswith(self.tmp))

    def tearDown(self):
        os.environ.pop("EXCEL_PATH", None)
        os.environ.pop("SETTINGS_PATH", None)
        _reload_stack()
        shutil.rmtree(self.tmp, ignore_errors=True)

    def _add(self, desc, amount, cur, cat="Прочее"):
        self.store.add_expense({
            "date": "2026-06-08", "description": desc, "amount": amount,
            "currency": cur, "category": cat, "amount_usd": amount,
        })

    def test_create_profile_makes_named_file(self):
        info = self.store.create_profile("Italy 2026")
        self.assertEqual(info["file"], "Italy 2026.xlsx")
        self.assertTrue(os.path.exists(os.path.join(self.tmp, "Italy 2026.xlsx")))
        self.assertEqual(self.settings.get_profile_file("Italy 2026"), "Italy 2026.xlsx")

    def test_profiles_are_isolated(self):
        # Дефолтный тур.
        self._add("Такси", 1000, "KZT")
        self.store.add_movement("Пополнение", "KZT", 5000)
        # Новый тур.
        self.store.create_profile("Italy 2026")
        self.settings.set_active_profile("Italy 2026")
        self._add("Pizza", 20, "USD")

        # Активный (Italy) видит только свою трату.
        totals = self.store.compute_totals()
        self.assertEqual(totals["count"], 1)
        self.assertEqual(totals["per_currency"]["USD"], 20)
        self.assertEqual(totals["per_currency"]["KZT"], 0)
        self.assertEqual(self.store.wallet_net()["KZT"], 0.0)

        # Дефолтный тур не затронут.
        self.settings.set_active_profile("FDTG tour 2026")
        totals = self.store.compute_totals()
        self.assertEqual(totals["count"], 1)
        self.assertEqual(totals["per_currency"]["KZT"], 1000)
        self.assertEqual(self.store.wallet_net()["KZT"], 5000.0)

    def test_import_creates_then_replaces_with_backup(self):
        name, replaced = self.store.import_profile_from_upload(
            "Greece 2026.xlsx", _upload_bytes([("Гирос", 10, "USD", "Питание", 10, 1)])
        )
        self.assertEqual(name, "Greece 2026")
        self.assertFalse(replaced)
        self.assertTrue(self.settings.profile_exists("Greece 2026"))
        self.assertEqual(self.settings.get_active_profile(), "Greece 2026")
        self.assertEqual(self.store.compute_totals()["count"], 1)

        # Повторный импорт того же тура — замена с бэкапом.
        name, replaced = self.store.import_profile_from_upload(
            "Greece 2026.xlsx",
            _upload_bytes([("Souvlaki", 12, "USD", "Питание", 12, 1),
                           ("Ouzo", 5, "USD", "Прочее", 5, 2)]),
        )
        self.assertTrue(replaced)
        self.assertEqual(self.store.compute_totals()["count"], 2)
        backups = glob.glob(os.path.join(self.tmp, "Greece 2026_backup_*.xlsx"))
        self.assertEqual(len(backups), 1)

    def test_migrate_legacy_single_workbook(self):
        # Старая схема: один файл с листами на тур + старый settings.json.
        wb = Workbook()
        ws = wb.active
        ws.title = "Расходы"
        ws.append(["Дата", "Описание", "Сумма", "Валюта", "Категория", "Сумма USD", "ID"])
        ws.append(["2026-06-08", "Отель", 100, "USD", "Отель", 100, 1])
        wb.create_sheet("Итоги")
        wsw = wb.create_sheet("Кошелёк")
        wsw.append(["Дата", "Тип", "Валюта", "Сумма", "Примечание"])
        wsw.append(["2026-06-08", "Старт", "USD", 500, "Начальный баланс"])
        wb.save(os.path.join(self.tmp, "expenses.xlsx"))

        with open(os.path.join(self.tmp, "settings.json"), "w", encoding="utf-8") as f:
            json.dump({
                "default_currency": "KZT",
                "profiles": {"FDTG tour 2026": {
                    "data": "Расходы", "summary": "Итоги", "wallet_sheet": "Кошелёк"}},
                "active_profile": "FDTG tour 2026",
            }, f, ensure_ascii=False)
        self.settings, self.store = _reload_stack()

        self.assertTrue(self.store.migrate_to_files())
        # Файл тура создан, данные на месте.
        self.assertTrue(os.path.exists(os.path.join(self.tmp, "FDTG tour 2026.xlsx")))
        self.assertEqual(self.settings.get_profile_file("FDTG tour 2026"),
                         "FDTG tour 2026.xlsx")
        self.assertEqual(self.store.compute_totals()["per_currency"]["USD"], 100)
        self.assertEqual(self.store.get_start_balance()["USD"], 500.0)
        # Старый общий файл выведен из обращения и забэкаплен.
        self.assertFalse(os.path.exists(os.path.join(self.tmp, "expenses.xlsx")))
        self.assertEqual(
            len(glob.glob(os.path.join(self.tmp, "expenses_legacy_backup_*.xlsx"))), 1)
        # Идемпотентность.
        self.assertFalse(self.store.migrate_to_files())


if __name__ == "__main__":
    unittest.main()
